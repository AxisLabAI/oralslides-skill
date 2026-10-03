#!/usr/bin/env python3
"""Local, resumable PowerPoint/PDF to narrated video workflow."""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any


DEFAULT_VOICES = {
    "zh": "zh-CN-XiaoxiaoNeural",
    "en": "en-US-JennyNeural",
    "ja": "ja-JP-NanamiNeural",
    "ko": "ko-KR-SunHiNeural",
    "es": "es-ES-ElviraNeural",
    "fr": "fr-FR-DeniseNeural",
    "de": "de-DE-KatjaNeural",
    "default": "en-US-JennyNeural",
}


class OralSlidesError(RuntimeError):
    pass


def run(command: list[str], timeout: int = 300) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        command,
        capture_output=True,
        text=True,
        timeout=timeout,
        check=False,
    )
    if result.returncode:
        detail = (result.stderr or result.stdout or "command failed").strip()
        raise OralSlidesError(f"{command[0]} failed: {detail[-2000:]}")
    return result


def find_command(*names: str) -> str | None:
    for name in names:
        found = shutil.which(name)
        if found:
            return found
    mac_soffice = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if "soffice" in names and mac_soffice.exists():
        return str(mac_soffice)
    return None


def doctor(_: argparse.Namespace) -> int:
    checks = {
        "LibreOffice": find_command("soffice", "libreoffice"),
        "pdftoppm": find_command("pdftoppm"),
        "ffmpeg": find_command("ffmpeg"),
        "ffprobe": find_command("ffprobe"),
    }
    failed = False
    for label, value in checks.items():
        print(f"{'OK' if value else 'MISSING':7} {label}: {value or '-'}")
        failed = failed or not value
    try:
        import edge_tts  # noqa: F401
        import pptx  # noqa: F401
        print("OK      Python packages: edge-tts, python-pptx")
    except ImportError as error:
        failed = True
        print(f"MISSING Python package: {error.name}")
    if failed:
        print("\nmacOS: brew install --cask libreoffice && brew install poppler ffmpeg")
        print("Ubuntu: sudo apt install libreoffice poppler-utils ffmpeg")
        print("Python: pip install -r requirements.txt")
        return 1
    return 0


def extract_pptx_text(source: Path) -> list[dict[str, Any]]:
    from pptx import Presentation

    presentation = Presentation(str(source))
    result: list[dict[str, Any]] = []
    for number, slide in enumerate(presentation.slides, start=1):
        blocks: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.strip()
                if text and text not in blocks:
                    blocks.append(text)
        result.append(
            {
                "slideNumber": number,
                "title": blocks[0] if blocks else "",
                "text": "\n".join(blocks),
                "bullets": blocks[1:] if len(blocks) > 1 else [],
            }
        )
    return result


def render_pdf(pdf: Path, slides_dir: Path) -> list[Path]:
    pdftoppm = find_command("pdftoppm")
    if not pdftoppm:
        raise OralSlidesError("pdftoppm is required; run doctor for installation help")
    prefix = slides_dir / "slide"
    run([pdftoppm, "-jpeg", "-r", "144", str(pdf), str(prefix)])
    images = sorted(slides_dir.glob("slide-*.jpg"))
    if not images:
        raise OralSlidesError("the deck did not contain any renderable pages")
    return images


def prepare(args: argparse.Namespace) -> int:
    source = Path(args.input).expanduser().resolve()
    if not source.exists():
        raise OralSlidesError(f"input does not exist: {source}")
    if source.suffix.lower() not in {".pptx", ".pdf"}:
        raise OralSlidesError("only .pptx and .pdf inputs are supported")
    work = Path(args.output).expanduser().resolve()
    slides_dir = work / "slides"
    slides_dir.mkdir(parents=True, exist_ok=True)
    extracted = extract_pptx_text(source) if source.suffix.lower() == ".pptx" else []

    with tempfile.TemporaryDirectory() as temp:
        if source.suffix.lower() == ".pptx":
            soffice = find_command("soffice", "libreoffice")
            if not soffice:
                raise OralSlidesError("LibreOffice is required; run doctor")
            run([soffice, "--headless", "--convert-to", "pdf", "--outdir", temp, str(source)])
            pdf = Path(temp) / f"{source.stem}.pdf"
            if not pdf.exists():
                candidates = list(Path(temp).glob("*.pdf"))
                if not candidates:
                    raise OralSlidesError("LibreOffice did not create a PDF")
                pdf = candidates[0]
        else:
            pdf = source
        rendered = render_pdf(pdf, slides_dir)

    slides: list[dict[str, Any]] = []
    for number, image in enumerate(rendered, start=1):
        stable_name = f"slide_{number:03d}.jpg"
        stable_path = slides_dir / stable_name
        if image != stable_path:
            image.replace(stable_path)
        metadata = extracted[number - 1] if number <= len(extracted) else {}
        slides.append(
            {
                "slideNumber": number,
                "image": f"slides/{stable_name}",
                "title": metadata.get("title", ""),
                "text": metadata.get("text", ""),
                "bullets": metadata.get("bullets", []),
            }
        )

    manifest = {"version": 1, "source": str(source), "slides": slides}
    (work / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    scripts_path = work / "scripts.json"
    if not scripts_path.exists():
        template = {
            "language": "auto",
            "voice": "auto",
            "rate": "+0%",
            "slides": [
                {"slideNumber": slide["slideNumber"], "script": ""} for slide in slides
            ],
        }
        scripts_path.write_text(json.dumps(template, ensure_ascii=False, indent=2) + "\n")
    print(f"Prepared {len(slides)} slides in {work}")
    return 0


def read_json(path: Path) -> dict[str, Any]:
    try:
        return json.loads(path.read_text())
    except (OSError, json.JSONDecodeError) as error:
        raise OralSlidesError(f"cannot read {path}: {error}") from error


def detect_language(text: str) -> str:
    if any("\u4e00" <= char <= "\u9fff" for char in text):
        return "zh"
    if any("\u3040" <= char <= "\u30ff" for char in text):
        return "ja"
    if any("\uac00" <= char <= "\ud7af" for char in text):
        return "ko"
    return "en"


def validate_scripts(work: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    manifest = read_json(work / "manifest.json")
    scripts_doc = read_json(work / "scripts.json")
    expected = [slide["slideNumber"] for slide in manifest.get("slides", [])]
    scripts = scripts_doc.get("slides", [])
    actual = [slide.get("slideNumber") for slide in scripts]
    if actual != expected:
        raise OralSlidesError(f"scripts.json slide numbers must exactly match {expected}")
    blanks = [str(slide["slideNumber"]) for slide in scripts if not str(slide.get("script", "")).strip()]
    if blanks:
        raise OralSlidesError(f"empty narration for slides: {', '.join(blanks)}")
    return scripts_doc, scripts


async def synthesize(work: Path, scripts_doc: dict[str, Any], scripts: list[dict[str, Any]], force: bool) -> None:
    import edge_tts

    audio_dir = work / "audio"
    audio_dir.mkdir(parents=True, exist_ok=True)
    state_path = audio_dir / "state.json"
    state = read_json(state_path) if state_path.exists() else {}
    next_state: dict[str, str] = {}
    all_text = "\n".join(str(item["script"]) for item in scripts)
    language = scripts_doc.get("language", "auto")
    if language == "auto":
        language = detect_language(all_text)
    voice = scripts_doc.get("voice", "auto")
    if voice == "auto":
        voice = DEFAULT_VOICES.get(language, DEFAULT_VOICES["default"])
    rate = scripts_doc.get("rate", "+0%")
    for item in scripts:
        number = int(item["slideNumber"])
        text = str(item["script"]).strip()
        fingerprint = hashlib.sha256(f"{voice}\0{rate}\0{text}".encode()).hexdigest()
        target = audio_dir / f"slide_{number:03d}.mp3"
        if not force and target.exists() and state.get(str(number)) == fingerprint:
            print(f"Reusing audio for slide {number}")
        else:
            print(f"Generating audio for slide {number}")
            await edge_tts.Communicate(text, voice, rate=rate).save(str(target))
        next_state[str(number)] = fingerprint
    state_path.write_text(json.dumps(next_state, indent=2) + "\n")


def tts(args: argparse.Namespace) -> int:
    work = Path(args.work_dir).expanduser().resolve()
    scripts_doc, scripts = validate_scripts(work)
    asyncio.run(synthesize(work, scripts_doc, scripts, args.force))
    return 0


def audio_duration(path: Path) -> float:
    ffprobe = find_command("ffprobe")
    if not ffprobe:
        raise OralSlidesError("ffprobe is required; run doctor")
    result = run([ffprobe, "-v", "error", "-show_entries", "format=duration", "-of", "default=nw=1:nk=1", str(path)])
    try:
        return max(float(result.stdout.strip()), 0.4)
    except ValueError as error:
        raise OralSlidesError(f"cannot determine audio duration for {path}") from error


def compose(args: argparse.Namespace) -> int:
    work = Path(args.work_dir).expanduser().resolve()
    manifest = read_json(work / "manifest.json")
    _, scripts = validate_scripts(work)
    ffmpeg = find_command("ffmpeg")
    if not ffmpeg:
        raise OralSlidesError("ffmpeg is required; run doctor")
    output = Path(args.output).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as temp:
        temp_dir = Path(temp)
        clips: list[Path] = []
        for slide, _script in zip(manifest["slides"], scripts, strict=True):
            number = int(slide["slideNumber"])
            image = work / slide["image"]
            audio = work / "audio" / f"slide_{number:03d}.mp3"
            if not image.exists() or not audio.exists():
                raise OralSlidesError(f"missing image or audio for slide {number}")
            clip = temp_dir / f"clip_{number:03d}.mp4"
            duration = audio_duration(audio)
            run([
                ffmpeg, "-y", "-loop", "1", "-i", str(image), "-i", str(audio),
                "-t", f"{duration:.3f}", "-r", "30", "-vf",
                "scale=1920:1080:force_original_aspect_ratio=decrease,pad=1920:1080:(ow-iw)/2:(oh-ih)/2:color=black,format=yuv420p",
                "-c:v", "libx264", "-tune", "stillimage", "-c:a", "aac",
                "-b:a", "192k", "-shortest", str(clip),
            ])
            clips.append(clip)
        concat = temp_dir / "concat.txt"
        concat.write_text("".join(f"file '{path}'\n" for path in clips))
        run([ffmpeg, "-y", "-f", "concat", "-safe", "0", "-i", str(concat), "-c:v", "libx264", "-c:a", "aac", "-movflags", "+faststart", str(output)], timeout=900)
    print(f"Created {output}")
    return 0


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser(prog="oralslides", description=__doc__)
    commands = root.add_subparsers(dest="command", required=True)
    command = commands.add_parser("doctor", help="check local dependencies")
    command.set_defaults(handler=doctor)
    command = commands.add_parser("prepare", help="render a PPTX/PDF and create script template")
    command.add_argument("input")
    command.add_argument("--output", "-o", default="oralslides-work")
    command.set_defaults(handler=prepare)
    command = commands.add_parser("tts", help="generate narration from scripts.json")
    command.add_argument("work_dir")
    command.add_argument("--force", action="store_true")
    command.set_defaults(handler=tts)
    command = commands.add_parser("compose", help="compose slide images and audio into MP4")
    command.add_argument("work_dir")
    command.add_argument("--output", "-o", default="oralslides.mp4")
    command.set_defaults(handler=compose)
    return root


def main() -> int:
    try:
        arguments = parser().parse_args()
        return int(arguments.handler(arguments))
    except (OralSlidesError, subprocess.TimeoutExpired) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
